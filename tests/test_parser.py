from pathlib import Path

from docx import Document

from qabank.parser import parse_docx


def build_fixture(path: Path) -> None:
    document = Document()
    document.add_paragraph("第一章", style="Heading 1")
    document.add_paragraph("单项选择题（共计1道）")
    document.add_paragraph("核心要义是（ ）。")
    document.add_paragraph("选项甲")
    document.add_paragraph("选项乙")
    document.add_paragraph("选项丙")
    document.add_paragraph("C")
    document.add_paragraph("多项选择题（共计1道）")
    document.add_paragraph("正确表述包括（ ）。", style="Heading 1")
    document.add_paragraph("表述甲", style="Heading 2")
    document.add_paragraph("表述乙", style="Heading 2")
    document.add_paragraph("表述丙", style="Heading 2")
    document.add_paragraph("表述丁", style="Heading 2")
    document.add_paragraph("【答案】ABD", style="Heading 3")
    document.add_paragraph("判断题（共计1道）")
    document.add_paragraph("这是一道判断题。", style="Heading 1")
    document.add_paragraph("正确", style="Heading 3")
    document.save(path)


def test_parser_keeps_one_question_per_record(tmp_path: Path):
    path = tmp_path / "fixture.docx"
    build_fixture(path)
    report = parse_docx(path, expected_count=3)

    assert report.valid
    assert report.summary()["type_counts"] == {
        "single_choice": 1,
        "multiple_choice": 1,
        "true_false": 1,
    }
    assert report.questions[0].options["C"] == "选项丙"
    assert report.questions[0].answer == "C"
    assert report.questions[1].answer == "ABD"
    assert report.questions[2].options == {}


def test_parser_rejects_answer_for_missing_option(tmp_path: Path):
    path = tmp_path / "broken.docx"
    document = Document()
    document.add_paragraph("单项选择题")
    document.add_paragraph("题干")
    document.add_paragraph("只有一个选项")
    document.add_paragraph("C")
    document.save(path)

    report = parse_docx(path, expected_count=None)
    assert not report.valid
    assert "不存在的选项" in report.issues[0].message
